#!/usr/bin/env python3
"""Import daily dish catalog from data/Dal_Birbante_jidla.xlsx into content.json."""

from __future__ import annotations

import json
import re
import unicodedata
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "Dal_Birbante_jidla.xlsx"
CONTENT = ROOT / "data" / "content.json"

CAT_MAP = {
    "polevky": "polevky",
    "panozzo": "panozzo",
    "testoviny": "pasta",
    "pasta": "pasta",
    "gnocchi": "gnocchi",
    "rizota": "rizota",
    "maso": "maso",
    "pizza": "pizza",
    "dezerty": "dezerty",
}

EMOJI = {
    "polevky": "🍅",
    "panozzo": "🥖",
    "pasta": "🍝",
    "gnocchi": "🧀",
    "rizota": "🍚",
    "maso": "🥩",
    "pizza": "🍕",
    "dezerty": "🍰",
    "ostatni": "",
}


def fold(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.lower()


def slugify(name: str) -> str:
    s = fold(name)
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:64] or "jidlo"


def map_cat(label: str) -> str:
    key = fold(label or "").strip()
    for prefix, cat in CAT_MAP.items():
        if key.startswith(prefix):
            return cat
    return "ostatni"


def split_note(desc: str) -> tuple[str, str]:
    desc = (desc or "").strip()
    m = re.search(r"\s*(\([^)]*(?:\+|Kč|kc)[^)]*\))\s*$", desc, re.I)
    if m:
        note = m.group(1).strip()
        body = desc[: m.start()].rstrip(" ,;")
        return body, note
    return desc, ""


def main() -> None:
    if not XLSX.exists():
        raise SystemExit(f"Missing {XLSX}")

    wb = load_workbook(XLSX, data_only=True)
    ws = wb.active
    used_ids: dict[str, int] = {}
    catalog = []
    for cat_label, name, desc, price in ws.iter_rows(min_row=2, values_only=True):
        if not name:
            continue
        category = map_cat(str(cat_label or ""))
        description, note = split_note(str(desc or ""))
        base = slugify(str(name))
        n = used_ids.get(base, 0)
        used_ids[base] = n + 1
        dish_id = base if n == 0 else f"{base}-{n + 1}"
        if isinstance(price, (int, float)):
            price_s = f"{int(price)} Kč"
        else:
            price_s = str(price or "").strip()
            if price_s and not price_s.endswith("Kč"):
                price_s = f"{price_s} Kč"
        catalog.append(
            {
                "id": dish_id,
                "name": str(name).strip().upper(),
                "price": price_s,
                "emoji": EMOJI.get(category, ""),
                "description": description,
                "note": note,
                "category": category,
            }
        )

    content = json.loads(CONTENT.read_text())
    old_catalog = content["daily"].get("catalog") or []
    old_today = content["daily"].get("todayIds") or []
    old_by_id = {d["id"]: d for d in old_catalog}
    new_by_fold: dict[str, list] = {}
    for d in catalog:
        new_by_fold.setdefault(fold(d["name"]), []).append(d)

    def best_match(old_id: str) -> str | None:
        old = old_by_id.get(old_id)
        if not old:
            return None
        for d in catalog:
            if d["id"] == old_id:
                return d["id"]
        old_name = fold(old["name"])
        if old_name in new_by_fold:
            return new_by_fold[old_name][0]["id"]
        candidates = []
        for d in catalog:
            fn = fold(d["name"])
            if old_name in fn or fn in old_name:
                # Prefer longer/closer name match (mozzarella variant over short)
                score = 0 if old_name == fn else abs(len(fn) - len(old_name))
                # Prefer when new name contains old (longer variant)
                if old_name in fn:
                    score -= 10
                candidates.append((score, -len(fn), d["id"]))
        if candidates:
            candidates.sort()
            return candidates[0][2]
        return None

    new_today = []
    for oid in old_today:
        mid = best_match(oid)
        if mid and mid not in new_today:
            new_today.append(mid)
        else:
            print("UNMATCHED today:", oid)

    by_id = {d["id"]: d for d in catalog}
    items = [{**by_id[i]} for i in new_today if i in by_id]
    content["daily"]["catalog"] = catalog
    content["daily"]["todayIds"] = new_today
    content["daily"]["items"] = items
    CONTENT.write_text(json.dumps(content, ensure_ascii=False, indent=2) + "\n")
    print("imported", len(catalog), "dishes")
    print(Counter(d["category"] for d in catalog))
    print("todayIds", new_today)


if __name__ == "__main__":
    main()
